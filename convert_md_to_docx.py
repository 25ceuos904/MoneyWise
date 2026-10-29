import re
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def build_docx(md_filepath, docx_filepath):
    doc = docx.Document()
    
    # Configure margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Palette
    PRIMARY_COLOR = RGBColor(10, 79, 64)    # #0a4f40 (Dark Green)
    SECONDARY_COLOR = RGBColor(15, 122, 98)  # #0f7a62 (Teal Green)
    TEXT_COLOR = RGBColor(16, 33, 28)       # #10211c
    MUTED_COLOR = RGBColor(93, 109, 103)    # #5d6d67

    # Normal Style
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Segoe UI'
    normal_style.font.size = Pt(10.5)
    normal_style.font.color.rgb = TEXT_COLOR

    with open(md_filepath, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    in_code_block = False
    code_lines = []
    in_table = False
    table_lines = []

    def process_inline(paragraph, text, bold_prefix=False):
        # Parses **bold**, `code`, etc.
        parts = re.split(r'(\*\*.*?\*\*|`.*?`|\*.*?\*)', text)
        for part in parts:
            if not part:
                continue
            if part.startswith('**') and part.endswith('**'):
                run = paragraph.add_run(part[2:-2])
                run.bold = True
            elif part.startswith('`') and part.endswith('`'):
                run = paragraph.add_run(part[1:-1])
                run.font.name = 'Consolas'
                run.font.size = Pt(9.5)
                run.font.color.rgb = PRIMARY_COLOR
            elif part.startswith('*') and part.endswith('*'):
                run = paragraph.add_run(part[1:-1])
                run.italic = True
            else:
                paragraph.add_run(part)

    def render_table(lines):
        # Filter out separator lines like |---|---|
        rows_data = []
        for l in lines:
            if re.match(r'^\s*\|?\s*:?-+:?\s*\|', l):
                continue
            cells = [c.strip() for c in l.strip('| \n').split('|')]
            if any(cells):
                rows_data.append(cells)

        if not rows_data:
            return

        cols_count = max(len(r) for r in rows_data)
        table = doc.add_table(rows=len(rows_data), cols=cols_count)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER

        for r_idx, row_data in enumerate(rows_data):
            row = table.rows[r_idx]
            is_header = (r_idx == 0)
            for c_idx, cell_value in enumerate(row_data):
                if c_idx < len(row.cells):
                    cell = row.cells[c_idx]
                    cell.text = ""
                    p = cell.paragraphs[0]
                    p.paragraph_format.space_before = Pt(4)
                    p.paragraph_format.space_after = Pt(4)
                    process_inline(p, cell_value)

                    set_cell_margins(cell, top=100, bottom=100, left=150, right=150)
                    if is_header:
                        set_cell_background(cell, "0f7a62")
                        for r in p.runs:
                            r.font.bold = True
                            r.font.color.rgb = RGBColor(255, 255, 255)
                    else:
                        if r_idx % 2 == 1:
                            set_cell_background(cell, "f4f7f5")
                        else:
                            set_cell_background(cell, "ffffff")

        p_space = doc.add_paragraph()
        p_space.paragraph_format.space_before = Pt(0)
        p_space.paragraph_format.space_after = Pt(4)

    i = 0
    while i < len(lines):
        line = lines[i]
        raw = line.rstrip('\r\n')

        # Handle Code Block
        if raw.startswith('```'):
            if in_code_block:
                # End code block
                in_code_block = False
                code_text = '\n'.join(code_lines)
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(6)
                p.paragraph_format.space_after = Pt(6)
                p.paragraph_format.left_indent = Inches(0.25)
                run = p.add_run(code_text)
                run.font.name = 'Consolas'
                run.font.size = Pt(9.5)
                run.font.color.rgb = PRIMARY_COLOR
                code_lines = []
            else:
                in_code_block = True
                code_lines = []
            i += 1
            continue

        if in_code_block:
            code_lines.append(raw)
            i += 1
            continue

        # Handle Table
        if raw.strip().startswith('|'):
            if not in_table:
                in_table = True
                table_lines = []
            table_lines.append(raw)
            i += 1
            continue
        else:
            if in_table:
                in_table = False
                render_table(table_lines)
                table_lines = []

        # Empty lines
        if not raw.strip():
            i += 1
            continue

        # Headings
        if raw.startswith('# '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(16)
            p.paragraph_format.space_after = Pt(8)
            run = p.add_run(raw[2:])
            run.font.size = Pt(22)
            run.font.bold = True
            run.font.color.rgb = PRIMARY_COLOR
        elif raw.startswith('## '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run(raw[3:])
            run.font.size = Pt(16)
            run.font.bold = True
            run.font.color.rgb = PRIMARY_COLOR
        elif raw.startswith('### '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(raw[4:])
            run.font.size = Pt(13)
            run.font.bold = True
            run.font.color.rgb = SECONDARY_COLOR
        elif raw.startswith('#### '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(raw[5:])
            run.font.size = Pt(11.5)
            run.font.bold = True
            run.font.color.rgb = SECONDARY_COLOR

        # Blockquote / Speech script
        elif raw.startswith('> '):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.4)
            p.paragraph_format.right_indent = Inches(0.4)
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            process_inline(p, raw[2:])

        # Bullet List
        elif raw.strip().startswith('- ') or raw.strip().startswith('* '):
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            content = raw.strip()[2:]
            process_inline(p, content)

        # Numbered List
        elif re.match(r'^\d+\.\s', raw.strip()):
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            content = re.sub(r'^\d+\.\s', '', raw.strip())
            process_inline(p, content)

        # Horizontal Rule
        elif raw.strip() in ['---', '***', '___']:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run('───────' * 8)
            run.font.color.rgb = MUTED_COLOR

        # Normal Paragraph
        else:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(4)
            process_inline(p, raw)

        i += 1

    if in_table:
        render_table(table_lines)

    doc.save(docx_filepath)
    print(f"Successfully generated {docx_filepath}")

if __name__ == '__main__':
    build_docx('MoneyWise_Study_Guide.md', 'MoneyWise_Study_Guide.docx')
